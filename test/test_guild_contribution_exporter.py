import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from function.scattered.guild_contribution_exporter import (
    export_selected_date_excel,
    get_rows_for_date,
)


MEMBERS_DATA = [
    {
        "name_image_hash": "alice",
        "data": {"2026-06-09": 3000, "2026-06-10": 3200},
        "data_week": {"2026-06-09": 1200, "2026-06-10": 2000},
    },
    {
        "name_image_hash": "bob",
        "data": {"2026-06-10": 800},
        "data_week": {"2026-06-10": 160},
    },
]


class TestGuildContributionExporter(unittest.TestCase):

    def test_rows_are_sorted_by_week_contribution(self):
        rows = get_rows_for_date(MEMBERS_DATA, "2026-06-10", Path("images"))
        self.assertEqual(
            [row["name_image_hash"] for row in rows], ["alice", "bob"])

    def test_export_creates_only_selected_date_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            output_dir = Path(directory)
            image_dir = output_dir / "guild_member_images"
            image_dir.mkdir()
            Image.new("RGB", (93, 35), "white").save(image_dir / "alice.png")
            output_path = export_selected_date_excel(
                MEMBERS_DATA, "2026-06-10", output_dir)

            self.assertTrue(output_path.is_file())
            self.assertEqual(output_path.name, "公会贡献数据导出_2026-06-10.xlsx")
            self.assertFalse(
                (output_dir / "guild_contribution_all_dates.xlsx").exists())

            workbook = load_workbook(output_path)
            self.assertEqual(workbook.sheetnames, ["2026-06-10"])
            worksheet = workbook["2026-06-10"]
            self.assertEqual(worksheet["B2"].value, 3200)
            self.assertEqual(worksheet["C2"].value, 2000)
            self.assertIsNone(worksheet["D1"].value)
            self.assertEqual(worksheet["A1"].font.name, "微软雅黑")
            self.assertEqual(worksheet["B2"].font.name, "微软雅黑")
            self.assertEqual(len(worksheet._images), 1)


if __name__ == "__main__":
    unittest.main()
