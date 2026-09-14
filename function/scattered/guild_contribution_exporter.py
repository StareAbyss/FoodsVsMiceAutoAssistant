"""公会贡献 Excel 导出工具。

原始实现由“小星蛋挞公会”的开发者“灼小星”提供，本模块在其源码基础上
按 FAA 的路径、异常处理和界面调用方式完成内置化。
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


NAME_IMAGE_WIDTH = 93
NAME_IMAGE_HEIGHT = 35
EXCEL_FONT_NAME = "微软雅黑"


def get_rows_for_date(
        members_data: list[dict], target_date: str, image_dir: Path) -> list[dict]:
    """整理指定日期的成员贡献数据，并按周贡献降序排列。

    Args:
        members_data: 公会管理器保存的成员数据。
        target_date: 要导出的日期，格式为 ``yyyy-MM-dd``。
        image_dir: 公会成员名称图片目录。

    Returns:
        可直接写入工作表的成员行数据。
    """
    rows = []
    for member in members_data:
        total_contribution = member.get("data", {}).get(target_date)
        week_contribution = member.get("data_week", {}).get(target_date)
        if total_contribution is None and week_contribution is None:
            continue

        name_image_hash = member.get("name_image_hash", "")
        rows.append({
            "name_image_hash": name_image_hash,
            "image_path": image_dir / f"{name_image_hash}.png",
            "total_contribution": int(total_contribution or 0),
            "week_contribution": int(week_contribution or 0),
        })

    rows.sort(key=lambda row: row["week_contribution"], reverse=True)
    return rows


def _safe_sheet_name(name: str) -> str:
    """清理 Excel 工作表名称中的非法字符。"""
    for char in ("\\", "/", "*", "?", ":", "[", "]"):
        name = name.replace(char, "_")
    return name[:31]


def _create_thin_border() -> Border:
    """创建导出表格统一使用的浅灰色细边框。"""
    side = Side(style="thin", color="CCCCCC")
    return Border(left=side, right=side, top=side, bottom=side)


def _setup_header(worksheet, headers: list[str]) -> Border:
    """写入工作表表头并返回数据行复用的边框。"""
    border = _create_thin_border()
    worksheet.append(headers)
    for column in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=column)
        cell.font = Font(name=EXCEL_FONT_NAME, bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    worksheet.freeze_panes = "A2"
    return border


def _setup_sheet_size(worksheet) -> None:
    """设置适合 93×35 成员名称图片的列宽。"""
    worksheet.column_dimensions["A"].width = 13
    for column in ("B", "C"):
        worksheet.column_dimensions[column].width = 10


def _add_member_image(worksheet, image_path: Path, cell: str) -> bool:
    """向单元格插入成员名称图片，图片缺失时返回 False。"""
    if not image_path.is_file():
        return False
    image = ExcelImage(str(image_path))
    image.width = NAME_IMAGE_WIDTH
    image.height = NAME_IMAGE_HEIGHT
    worksheet.add_image(image, cell)
    return True


def _style_data_row(
        worksheet, row: int, start_column: int, end_column: int,
        border: Border) -> None:
    """为一段数据行统一设置微软雅黑字体、居中和边框。"""
    for column in range(start_column, end_column + 1):
        cell = worksheet.cell(row=row, column=column)
        cell.font = Font(name=EXCEL_FONT_NAME)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def export_selected_date_excel(
        members_data: list[dict], target_date: str, guild_manager_dir: Path) -> Path:
    """导出指定日期的公会贡献表。

    本功能原始源码由“小星蛋挞公会”的开发者“灼小星”提供，内置时按 FAA
    的路径、异常处理和界面调用方式进行了适配。

    Args:
        members_data: 公会管理器保存的成员数据。
        target_date: 要导出的日期，格式为 ``yyyy-MM-dd``。
        guild_manager_dir: 公会管理器的数据及导出目录。

    Returns:
        生成的 Excel 文件路径。

    Raises:
        ValueError: 指定日期没有贡献数据。
    """
    image_dir = guild_manager_dir / "guild_member_images"
    rows = get_rows_for_date(members_data, target_date, image_dir)
    if not rows:
        raise ValueError(f"所选日期 {target_date} 没有可导出的贡献数据")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = _safe_sheet_name(target_date)
    border = _setup_header(worksheet, ["成员名", "总贡献", "周贡献"])

    for row_index, row_data in enumerate(rows, start=2):
        if not _add_member_image(worksheet, row_data["image_path"], f"A{row_index}"):
            worksheet.cell(row=row_index, column=1).value = (
                f"图片不存在: {row_data['image_path'].name}")
        worksheet.cell(row=row_index, column=2).value = row_data["total_contribution"]
        worksheet.cell(row=row_index, column=3).value = row_data["week_contribution"]
        worksheet.row_dimensions[row_index].height = 32
        _style_data_row(worksheet, row_index, 1, 3, border)

    _setup_sheet_size(worksheet)
    worksheet.auto_filter.ref = f"A1:C{len(rows) + 1}"
    guild_manager_dir.mkdir(parents=True, exist_ok=True)
    output_path = guild_manager_dir / f"公会贡献数据导出_{target_date}.xlsx"
    workbook.save(output_path)
    return output_path
